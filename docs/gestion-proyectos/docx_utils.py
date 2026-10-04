# -*- coding: utf-8 -*-
"""Utilidades para construir el informe en Word con python-docx."""
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1F, 0x4E, 0x79)
AZUL_HEX = "1F4E79"
GRIS_HEX = "F2F5F9"
FUENTE = "Calibri"

_contador = {"tabla": 0, "figura": 0}


def nuevo_documento():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    est = doc.styles
    n = est["Normal"]
    n.font.name = FUENTE
    n.font.size = Pt(10.5)
    n.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.12
    for nombre, tam, color, antes, despues in (("Heading 1", 18, AZUL, 18, 8), ("Heading 2", 14, AZUL, 14, 6), ("Heading 3", 12, RGBColor(0x2E, 0x75, 0xB6), 10, 4)):
        h = est[nombre]
        h.font.name = FUENTE
        h.font.size = Pt(tam)
        h.font.bold = True
        h.font.color.rgb = color
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
        h.element.rPr.rFonts.set(qn("w:ascii"), FUENTE)
        h.element.rPr.rFonts.set(qn("w:hAnsi"), FUENTE)
        h.paragraph_format.space_before = Pt(antes)
        h.paragraph_format.space_after = Pt(despues)
        h.paragraph_format.keep_with_next = True
    for nombre in ("List Bullet", "List Number"):
        est[nombre].font.name = FUENTE
        est[nombre].font.size = Pt(10.5)
        est[nombre].paragraph_format.space_after = Pt(3)
    # actualizar campos (índice) al abrir
    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)
    return doc


def _campo(par, instruccion, texto_previo=""):
    run = par.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instruccion
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t"); t.text = texto_previo
    f3 = OxmlElement("w:fldChar"); f3.set(qn("w:fldCharType"), "end")
    for el in (f1, it, f2, t, f3):
        run._r.append(el)
    return run


def pie_y_encabezado(doc, titulo_corto):
    for sec in doc.sections:
        sec.header.is_linked_to_previous = False
        sec.footer.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]
        for r in list(hp.runs):
            r._r.getparent().remove(r._r)
        hp.text = ""
        r = hp.add_run("Universidad Nacional de Cajamarca · EPIS · Gestión de Proyectos de Sistemas I · 2026-I")
        r.font.size = Pt(8); r.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp = sec.footer.paragraphs[0]
        for r_ in list(fp.runs):
            r_._r.getparent().remove(r_._r)
        fp.text = ""
        r = fp.add_run(titulo_corto + "  ·  Página ")
        r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
        c = _campo(fp, "PAGE", "1"); c.font.size = Pt(8.5)
        r = fp.add_run(" de "); r.font.size = Pt(8.5); r.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
        c = _campo(fp, "NUMPAGES", "1"); c.font.size = Pt(8.5)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def salto_pagina(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def seccion(doc, horizontal=False):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if horizontal:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(1.8)
        s.top_margin, s.bottom_margin = Cm(2.0), Cm(1.8)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin, s.bottom_margin = Cm(2.2), Cm(2.0)
    return s


def ancho_util(doc):
    s = doc.sections[-1]
    return (s.page_width - s.left_margin - s.right_margin) / 360000.0   # EMU -> cm


def h1(doc, texto):
    return doc.add_heading(texto, level=1)


def h2(doc, texto):
    return doc.add_heading(texto, level=2)


def h3(doc, texto):
    return doc.add_heading(texto, level=3)


def parrafo(doc, texto, negrita=False, cursiva=False, alineacion=None, tam=None, color=None, despues=None):
    p = doc.add_paragraph()
    _runs(p, texto, negrita, cursiva, tam, color)
    if alineacion == "centro":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif alineacion == "justificado":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if despues is not None:
        p.paragraph_format.space_after = Pt(despues)
    return p


def _runs(p, texto, negrita=False, cursiva=False, tam=None, color=None):
    """Soporta **negrita** y *cursiva* dentro del texto."""
    partes = texto.split("**")
    for i, parte in enumerate(partes):
        if not parte:
            continue
        sub = parte.split("*")
        for j, trozo in enumerate(sub):
            if not trozo:
                continue
            r = p.add_run(trozo)
            r.bold = negrita or (i % 2 == 1)
            r.italic = cursiva or (j % 2 == 1)
            if tam:
                r.font.size = Pt(tam)
            if color:
                r.font.color.rgb = color


def viñetas(doc, items, estilo="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=estilo)
        _runs(p, it)


def numerada(doc, items):
    viñetas(doc, items, "List Number")


def _sombrear(celda, hex_color):
    tcPr = celda._tc.get_or_add_tcPr()
    for e in tcPr.findall(qn("w:shd")):
        tcPr.remove(e)
    shd = parse_xml(r'<w:shd {} w:val="clear" w:color="auto" w:fill="{}"/>'.format(nsdecls("w"), hex_color))
    tcPr.append(shd)


def _bordes(tabla, color="BFC9D4"):
    tbl = tabla._tbl
    tblPr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def _margenes_celda(tabla, arriba=40, abajo=40, izq=70, der=70):
    tblPr = tabla._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for lado, v in (("top", arriba), ("left", izq), ("bottom", abajo), ("right", der)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa")
        m.append(e)
    tblPr.append(m)


def tabla(doc, encabezados, filas, anchos=None, tam=8.5, alinear=None, zebra=True, primera_negrita=False,
          titulo=None, color_encabezado=AZUL_HEX, resaltar=None):
    """Crea una tabla con encabezado repetible. `anchos` en cm (se escala al ancho útil).
    `alinear`: lista con 'l','c','r' por columna. `resaltar`: {indice_fila: hex} para sombrear filas."""
    if titulo:
        _contador["tabla"] += 1
        pc = doc.add_paragraph()
        r = pc.add_run(f"Tabla {_contador['tabla']}. "); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = AZUL
        r2 = pc.add_run(titulo); r2.font.size = Pt(9.5); r2.italic = True
        pc.paragraph_format.keep_with_next = True
        pc.paragraph_format.space_after = Pt(3)
    n = len(encabezados)
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    _bordes(t)
    _margenes_celda(t)
    util = ancho_util(doc)
    if anchos is None:
        anchos = [util / n] * n
    esc = util / sum(anchos)
    anchos = [a * esc for a in anchos]
    # encabezado
    for j, h in enumerate(encabezados):
        c = t.rows[0].cells[j]
        c.width = Cm(anchos[j])
        _sombrear(c, color_encabezado)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(h)); r.bold = True; r.font.size = Pt(tam); r.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for i, fila in enumerate(filas):
        row = t.add_row()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true")
        row._tr.get_or_add_trPr().append(cs)
        for j in range(n):
            c = row.cells[j]
            c.width = Cm(anchos[j])
            val = fila[j] if j < len(fila) else ""
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            texto = "" if val is None else str(val)
            lineas = texto.split("\n")
            for k, ln in enumerate(lineas):
                if k > 0:
                    p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
                _runs(p, ln, negrita=(primera_negrita and j == 0), tam=tam)
                if alinear and alinear[j] == "c":
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif alinear and alinear[j] == "r":
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if resaltar and i in resaltar:
                _sombrear(c, resaltar[i])
            elif zebra and i % 2 == 1:
                _sombrear(c, GRIS_HEX)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def figura(doc, ruta, ancho_cm, titulo):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(ruta, width=Cm(ancho_cm))
    _contador["figura"] += 1
    pc = doc.add_paragraph()
    pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = pc.add_run(f"Figura {_contador['figura']}. "); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = AZUL
    r2 = pc.add_run(titulo); r2.font.size = Pt(9.5); r2.italic = True


def caja(doc, texto, color="EAF2FB", titulo=None):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _bordes(t, "9DB7D5")
    _margenes_celda(t, 80, 80, 140, 140)
    c = t.rows[0].cells[0]
    c.width = Cm(ancho_util(doc))
    _sombrear(c, color)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    if titulo:
        r = p.add_run(titulo); r.bold = True; r.font.color.rgb = AZUL; r.font.size = Pt(10)
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(0)
    _runs(p, texto, tam=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def indice(doc):
    p = doc.add_paragraph()
    _campo(p, 'TOC \\o "1-3" \\h \\z \\u', "Haga clic derecho aquí y elija «Actualizar campo» (o presione F9) para generar el índice.")


def firmas(doc, firmantes):
    """Bloque de firmas: lista de (nombre, cargo)."""
    n = len(firmantes)
    t = doc.add_table(rows=2, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (nombre, cargo) in enumerate(firmantes):
        c = t.rows[0].cells[j]
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("\n\n_______________________________")
        c2 = t.rows[1].cells[j]
        p2 = c2.paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run(nombre + "\n"); r.bold = True; r.font.size = Pt(9)
        r = p2.add_run(cargo); r.font.size = Pt(8.5)
    doc.add_paragraph()
